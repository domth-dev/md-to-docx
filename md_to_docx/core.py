from typing import List, Optional
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from .parser import MarkdownParser
from .styles import StyleManager
from .numbering import NumberingManager
from .math import MathEngine
from .renderers import BlockRenderer, InlineRenderer
from .net import HttpClient

class MarkdownToDocx:
    def __init__(self,
                 template_path: Optional[str] = None,
                 apply_base_font: bool = False,
                 base_font: str = "Calibri",
                 base_size_pt: int = 11,
                 strict_math: bool = False,
                 http_download: bool = True):
        self.doc = Document(template_path) if template_path else Document()

        # services
        self.parser = MarkdownParser()
        self.styles = StyleManager(self.doc, base_font, base_size_pt, apply_base_font)
        self.styles.ensure()

        self.numbering = NumberingManager(self.doc)
        self.numbering.initialize()

        self.math = MathEngine(strict=strict_math)
        self.http = HttpClient() if http_download else None

        self.blocks = BlockRenderer(self.doc, self.numbering, self.math)
        self.inline = InlineRenderer(self.doc, self.math, http_client=self.http)

    def convert_text(self, markdown_text: str, output_path: Optional[str] = None) -> Document:
        tokens = self.parser.parse(markdown_text)
        self._render(tokens)
        if output_path:
            self.doc.save(output_path)
        return self.doc

    # --- render loop ---
    def _render(self, tokens):
        list_stack: List[dict] = []
        i = 0
        while i < len(tokens):
            t = tokens[i]

            if t.type == "heading_open":
                level = int(t.tag[1])
                text = self.blocks.collect_heading_text(tokens, i + 1)
                self.doc.add_heading(text, level=level)
                i = self._find_close(tokens, i, "heading_close")

            elif t.type == "paragraph_open":
                p = self.doc.add_paragraph()
                i = self._render_inline_para(tokens, i + 1, p)

            elif t.type == "blockquote_open":
                i += 1
                while tokens[i].type != "blockquote_close":
                    if tokens[i].type == "paragraph_open":
                        p = self.doc.add_paragraph()
                        try: p.style = self.doc.styles["Intense Quote"]
                        except Exception: pass
                        i = self._render_inline_para(tokens, i + 1, p)
                    elif tokens[i].type == "heading_open":
                        level = int(tokens[i].tag[1])
                        text = self.blocks.collect_heading_text(tokens, i + 1)
                        self.doc.add_heading(text, level=level)
                        i = self._find_close(tokens, i, "heading_close")
                    else:
                        i += 1
                i += 1

            elif t.type in ("bullet_list_open", "ordered_list_open"):
                list_stack.append({"type": t.type, "level": len(list_stack)})
                i += 1

            elif t.type == "list_item_open":
                level = len(list_stack) - 1
                is_ordered = (list_stack[-1]["type"] == "ordered_list_open")
                cur_p = self.doc.add_paragraph()
                self.numbering.apply(cur_p, ordered=is_ordered, level=level)

                i += 1
                while tokens[i].type != "list_item_close":
                    if tokens[i].type == "paragraph_open":
                        if cur_p.text or cur_p.runs:
                            cur_p = self.doc.add_paragraph()
                            self.numbering.apply(cur_p, ordered=is_ordered, level=level)
                        i = self._render_inline_para(tokens, i + 1, cur_p, append=True)
                    elif tokens[i].type == "inline":
                        self.inline.render(cur_p, tokens[i].children or [])
                        i += 1
                    elif tokens[i].type in ("bullet_list_open", "ordered_list_open"):
                        list_stack.append({"type": tokens[i].type, "level": len(list_stack)}); i += 1
                    elif tokens[i].type == "list_item_open":
                        level = len(list_stack) - 1
                        is_ord = (list_stack[-1]["type"] == "ordered_list_open")
                        cur_p = self.doc.add_paragraph()
                        self.numbering.apply(cur_p, ordered=is_ord, level=level)
                        i += 1
                    elif tokens[i].type in ("bullet_list_close", "ordered_list_close"):
                        list_stack.pop(); i += 1
                    else:
                        i += 1
                i += 1

            elif t.type in ("bullet_list_close", "ordered_list_close"):
                if list_stack: list_stack.pop()
                i += 1

            elif t.type == "fence":
                p = self.doc.add_paragraph(); p.style = self.doc.styles["Code Block"]
                p.add_run(t.content.rstrip("\n"))
                self.blocks.shade_paragraph(p); i += 1

            elif t.type == "code_block":
                p = self.doc.add_paragraph(); p.style = self.doc.styles["Code Block"]
                p.add_run(t.content.rstrip("\n"))
                self.blocks.shade_paragraph(p); i += 1

            elif t.type == "hr":
                self.blocks.add_horizontal_rule(self.doc); i += 1

            elif t.type == "table_open":
                i = self.blocks.render_table(tokens, i)

            elif t.type == "math_block":
                p = self.doc.add_paragraph()
                p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                self.math.add_math(p, t.content, display=True); i += 1

            else:
                i += 1

    def _render_inline_para(self, tokens, start, paragraph, append=False):
        i = start
        while tokens[i].type != "paragraph_close":
            if tokens[i].type == "inline":
                self.inline.render(paragraph, tokens[i].children or [])
                i += 1
            else:
                i += 1
        return i + 1

    def _find_close(self, tokens, start, close_type):
        i = start
        while i < len(tokens) and tokens[i].type != close_type:
            i += 1
        return i + 1