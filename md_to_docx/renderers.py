import os
from typing import List
from docx.shared import Inches
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

class LinkHelper:
    @staticmethod
    def add_hyperlink(paragraph, text, url):
        part = paragraph.part
        r_id = part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("r:id"), r_id)
        run = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
        color = OxmlElement("w:color"); color.set(qn("w:val"), "0000FF"); rPr.append(color)
        t = OxmlElement("w:t"); t.text = text
        run.append(rPr); run.append(t); hyperlink.append(run)
        paragraph._p.append(hyperlink)

class BlockRenderer:
    def __init__(self, doc, numbering_manager, math_engine):
        self.doc = doc
        self.num = numbering_manager
        self.math = math_engine

    @staticmethod
    def shade_paragraph(paragraph):
        pPr = paragraph._p.get_or_add_pPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "EDEDED"); pPr.append(shd)

    @staticmethod
    def add_horizontal_rule(doc):
        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "auto")
        pBdr.append(bottom); pPr.append(pBdr)

    @staticmethod
    def collect_heading_text(tokens, start_index) -> str:
        text = []; i = start_index
        while tokens[i].type != "heading_close":
            if tokens[i].type == "inline":
                for ch in (tokens[i].children or []):
                    if ch.type in ("text","code_inline"): text.append(ch.content)
                    elif ch.type in ("softbreak","hardbreak"): text.append("\n")
            i += 1
        return "".join(text).strip()

    @staticmethod
    def inline_plain(children) -> str:
        out = []
        for ch in children:
            if ch.type in ("text","code_inline"): out.append(ch.content)
            elif ch.type in ("softbreak","hardbreak"): out.append("\n")
        return "".join(out)

    def render_table(self, tokens, start_index) -> int:
        doc = self.doc
        i = start_index; headers, rows = [], []

        def collect_row(j, expected_close):
            cells, aligns = [], []
            while tokens[j].type != expected_close:
                if tokens[j].type in ("th_open","td_open"):
                    align = None
                    if tokens[j].attrs:
                        style = tokens[j].attrs.get("style") or ""
                        if "text-align:center" in style: align = "center"
                        elif "text-align:right" in style: align = "right"
                        elif "text-align:left"  in style: align = "left"
                    j += 1
                    parts = []
                    while tokens[j].type not in ("th_close","td_close"):
                        if tokens[j].type == "inline":
                            parts.append(self.inline_plain(tokens[j].children or []))
                        j += 1
                    cells.append("".join(parts).strip()); aligns.append(align)
                j += 1
            return cells, aligns

        i += 1; header_aligns = []
        while tokens[i].type != "table_close":
            if tokens[i].type == "thead_open":
                i += 1
                while tokens[i].type != "thead_close":
                    if tokens[i].type == "tr_open":
                        i += 1
                        row, header_aligns = collect_row(i, "tr_close")
                        headers.append(row)
                        while tokens[i].type != "tr_close": i += 1
                        i += 1
                    else: i += 1
                i += 1
            elif tokens[i].type == "tbody_open":
                i += 1
                while tokens[i].type != "tbody_close":
                    if tokens[i].type == "tr_open":
                        i += 1
                        row, aligns = collect_row(i, "tr_close")
                        rows.append((row, aligns))
                        while tokens[i].type != "tr_close": i += 1
                        i += 1
                    else: i += 1
                i += 1
            else: i += 1

        n_cols = len(headers[0]) if headers else (len(rows[0][0]) if rows else 0)
        if n_cols == 0: return i + 1

        n_rows = (1 if headers else 0) + len(rows)
        table = doc.add_table(rows=n_rows, cols=n_cols)
        if "Table Grid" in doc.styles: table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        r_index = 0
        if headers:
            for c, text in enumerate(headers[0]):
                cell = table.cell(0, c); cell.text = text
                for p in cell.paragraphs:
                    for run in p.runs: run.bold = True
                    if header_aligns and c < len(header_aligns) and header_aligns[c]:
                        p.alignment = {
                            "center": WD_ALIGN_PARAGRAPH.CENTER,
                            "right":  WD_ALIGN_PARAGRAPH.RIGHT,
                            "left":   WD_ALIGN_PARAGRAPH.LEFT,
                        }[header_aligns[c]]
            r_index = 1

        for r, (row, aligns) in enumerate(rows):
            for c in range(n_cols):
                cell = table.cell(r_index + r, c); cell.text = row[c] if c < len(row) else ""
                for p in cell.paragraphs:
                    if aligns and c < len(aligns) and aligns[c]:
                        p.alignment = {
                            "center": WD_ALIGN_PARAGRAPH.CENTER,
                            "right":  WD_ALIGN_PARAGRAPH.RIGHT,
                            "left":   WD_ALIGN_PARAGRAPH.LEFT,
                        }[aligns[c]]

        doc.add_paragraph()
        return i + 1

class InlineRenderer:
    def __init__(self, doc, math_engine, http_client=None):
        self.doc = doc
        self.math = math_engine
        self.http = http_client  # optional

    def render(self, paragraph, children):
        bold = italic = strike = code = False
        link_stack: List[str] = []
        i = 0
        while i < len(children):
            ch = children[i]
            if ch.type == "text":
                run = paragraph.add_run(ch.content)
                if code: run.style = self.doc.styles["Inline Code"]
                run.bold = True if bold else None
                run.italic = True if italic else None
                run.font.strike = True if strike else None

            elif ch.type in ("softbreak","hardbreak"):
                paragraph.add_run("\n")

            elif ch.type == "code_inline":
                run = paragraph.add_run(ch.content)
                run.style = self.doc.styles["Inline Code"]

            elif ch.type == "strong_open": bold = True
            elif ch.type == "strong_close": bold = False
            elif ch.type == "em_open": italic = True
            elif ch.type == "em_close": italic = False
            elif ch.type == "s_open": strike = True
            elif ch.type == "s_close": strike = False

            elif ch.type == "link_open":
                href = ""
                if ch.attrs:
                    for k,v in ch.attrs.items():
                        if k == "href": href = v
                link_stack.append(href)
            elif ch.type == "link_close":
                if link_stack: link_stack.pop()

            elif ch.type == "image":
                src = ch.attrs.get("src","") if ch.attrs else ""
                alt = ch.attrs.get("alt","") if ch.attrs else ""
                if src.startswith(("http://","https://")) and self.http:
                    try:
                        stream = self.http.fetch_image(src)
                        run = paragraph.add_run(); run.add_picture(stream, width=Inches(4))
                        paragraph = self.doc.add_paragraph()
                    except Exception:
                        LinkHelper.add_hyperlink(paragraph, alt or os.path.basename(src) or src, src)
                elif src and os.path.exists(src):
                    try:
                        paragraph.add_run("\n")
                        self.doc.add_picture(src, width=Inches(4))
                        paragraph = self.doc.add_paragraph()
                    except Exception:
                        LinkHelper.add_hyperlink(paragraph, alt or os.path.basename(src), src)
                else:
                    if src:
                        LinkHelper.add_hyperlink(paragraph, alt or os.path.basename(src), src)
                    else:
                        paragraph.add_run(alt or "[Bild]")

            elif ch.type == "math_inline":
                self.math.add_math(paragraph, ch.content, display=False)

            if link_stack and ch.type == "text":
                last_text = ch.content
                paragraph.runs[-1].text = ""
                LinkHelper.add_hyperlink(paragraph, last_text, link_stack[-1])

            i += 1
